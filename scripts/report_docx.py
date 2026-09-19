"""Build a fixed-layout scientific DOCX from report_data.json."""

from __future__ import annotations

import json
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = "1F4E78"
LIGHT_BLUE = "D9E2F3"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


LABELS = {
    "analysis_level": "Analysis level",
    "assay_type": "Experiment type",
    "direction": "Curve direction",
    "concentration_unit": "Concentration unit",
    "response_unit": "Response unit",
    "standard_replicate_policy": "Standard repeat policy",
    "curve_validated": "Curve validated",
    "validated_quantification_enabled": "Validated quantification enabled",
    "sample_id": "Sample identifier",
    "replicate_type": "Repeat type",
    "source_location": "Source location",
    "record_number": "Record number",
    "n_measurements": "Measurement rows",
    "mean": "Mean",
    "median": "Median",
    "min": "Minimum",
    "max": "Maximum",
    "sample_sd": "Sample SD",
    "group": "Group",
    "independent_biological_n": "Independent biological n",
    "experimental_unit_id": "Experimental unit",
    "aggregated_value": "Unit value",
    "technical_repeat_count": "Technical repeat count",
    "mean_difference": "Mean difference",
    "mean_difference_ci_lower": "95% CI lower",
    "mean_difference_ci_upper": "95% CI upper",
    "t_statistic": "t statistic",
    "degrees_of_freedom": "Degrees of freedom",
    "p_value": "Two-sided p value",
    "alpha": "Alpha",
    "alternative": "Alternative",
    "method": "Method",
    "lower_asymptote": "Lower asymptote L",
    "upper_asymptote": "Upper asymptote U",
    "midpoint_concentration": "Midpoint concentration C",
    "slope_magnitude": "Slope magnitude B",
    "n_fit_levels": "Fitted concentration levels",
    "n_raw_standard_measurements": "Raw standard measurements",
    "observed_standard_concentration_span": "Observed positive concentration span",
    "sse": "SSE",
    "rmse": "RMSE",
    "r_squared": "Descriptive R squared",
    "jacobian_rank": "Jacobian rank",
    "jacobian_condition_number": "Jacobian condition number",
    "midpoint_within_observed_positive_span": "Midpoint within observed positive span",
    "concentration": "Concentration",
    "observed": "Observed response",
    "predicted": "Predicted response",
    "residual": "Residual",
    "status": "Status",
    "response": "Response",
    "dilution_factor": "Dilution factor",
    "concentration_in_assayed_sample": "Assayed sample concentration",
    "concentration_in_original_sample": "Original sample concentration",
    "reason": "Reason",
}


def label(key: str) -> str:
    return LABELS.get(key, key.replace("_", " ").capitalize())


def display(value: Any) -> str:
    if value is None:
        return "Not calculated"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return format(value, ".10g")
    if isinstance(value, list | tuple):
        return ", ".join(display(item) for item in value)
    if isinstance(value, dict):
        return "; ".join(f"{label(str(k))}: {display(v)}" for k, v in value.items())
    return str(value)


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell: Any, color: str = BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_cell_padding(cell: Any, twips: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(twips))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table(doc: Document, headers: list[str], rows: list[list[Any]]) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_padding(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = display(value)
            set_cell_borders(cell)
            set_cell_padding(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = BLACK


def add_key_value_table(doc: Document, values: dict[str, Any]) -> None:
    add_table(doc, ["Field", "Value"], [[label(key), value] for key, value in values.items()])


def add_issue_table(doc: Document, issues: list[dict[str, Any]]) -> None:
    if not issues:
        doc.add_paragraph("No warnings or issues were recorded in the upstream artifacts.")
        return
    rows = [
        [
            item.get("code"),
            item.get("severity"),
            item.get("message"),
            item.get("location"),
            item.get("suggested_action"),
        ]
        for item in issues
    ]
    add_table(doc, ["Code", "Severity", "Message", "Location", "Recommended action"], rows)


def add_chart(doc: Document, chart_dir: Path, relative: str) -> None:
    path = chart_dir.parent / relative
    if not path.is_file():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.35))
    caption = doc.add_paragraph(relative.replace("_", " ").replace("/", " "))
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = doc.styles
    style_sizes = (("Normal", 10.5), ("Title", 20), ("Heading 1", 14), ("Heading 2", 11.5))
    for style_name, size in style_sizes:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
    styles["Title"].font.bold = True
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.bold = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("BioLab Copilot | Deterministic scientific report")
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(89, 89, 89)


def write_report(data_path: Path, chart_dir: Path, output_path: Path) -> None:
    data = json.loads(data_path.read_text(encoding="utf-8"))
    doc = Document()
    style_document(doc)
    properties = doc.core_properties
    properties.author = "BioLab Copilot"
    properties.last_modified_by = "BioLab Copilot"
    properties.title = str(data["title"])
    properties.subject = str(data["report_type"])
    fixed_timestamp = datetime(2000, 1, 1, tzinfo=UTC)
    properties.created = fixed_timestamp
    properties.modified = fixed_timestamp

    title = doc.add_paragraph(style="Title")
    title.add_run(str(data["title"]))
    subtitle = doc.add_paragraph()
    subtitle.add_run(f"Report type: {data['report_type']} | Report ID: {data['report_id']}")
    subtitle.runs[0].font.color.rgb = RGBColor(89, 89, 89)

    add_heading(doc, "Purpose and Scope")
    doc.add_paragraph(str(data["scope"]))
    if data["report_type"] == "elisa_4pl" and data["results"].get("inverse") is not None:
        doc.add_paragraph("Research-only; curve and quantitative range not validated.")

    add_heading(doc, "Data and Analysis Sources")
    source = data["source"]
    add_key_value_table(
        doc,
        {
            "source_filename": source.get("source_filename"),
            "source_sha256": source.get("source_sha256"),
            "input_artifact_sha256": source.get("input_artifact_sha256"),
            "upstream_run_id": source.get("upstream_run_id"),
            "upstream_manifest_sha256": source.get("upstream_manifest_sha256"),
        },
    )

    add_heading(doc, "QC and Warning Summary")
    add_issue_table(doc, list(data["qc"]["issues"]))

    add_heading(doc, "Confirmed Methods")
    add_key_value_table(doc, dict(data["methods"]))

    add_heading(doc, "Results")
    results = data["results"]
    report_type = data["report_type"]
    if report_type == "generic_grouped":
        stats = results["group_statistics"]
        add_table(
            doc,
            ["Group", "Measurement rows", "Mean", "Median", "Minimum", "Maximum", "Sample SD"],
            [
                [
                    item.get("group"),
                    item.get("n_measurements"),
                    item.get("mean"),
                    item.get("median"),
                    item.get("min"),
                    item.get("max"),
                    item.get("sample_sd"),
                ]
                for item in stats
            ],
        )
        doc.add_paragraph("Independent biological n: Not calculated.")
        measurements = results["measurements"]
        add_table(
            doc,
            [
                "Record number",
                "Sample identifier",
                "Group",
                "Measurement",
                "Repeat type",
                "Source location",
            ],
            [
                [
                    item.get("record_number"),
                    item.get("sample_id"),
                    item.get("group"),
                    item.get("measurement"),
                    item.get("replicate_type"),
                    item.get("source_location"),
                ]
                for item in measurements
            ],
        )
    elif report_type == "welch_two_group":
        stats = results["group_statistics"]
        add_table(
            doc,
            ["Group", "Measurement rows", "Experimental units", "Mean", "Sample SD"],
            [
                [
                    item.get("group"),
                    item.get("n_measurements"),
                    item.get("n_experimental_units"),
                    item.get("mean"),
                    item.get("sample_sd"),
                ]
                for item in stats
            ],
        )
        comparison = results["comparison"]
        add_key_value_table(doc, comparison)
        add_heading(doc, "Experimental Units", level=2)
        add_table(
            doc,
            [
                "Experimental unit",
                "Group",
                "Unit value",
                "Measurement rows",
                "Technical repeat count",
                "Source records",
            ],
            [
                [
                    item.get("experimental_unit_id"),
                    item.get("group"),
                    item.get("aggregated_value"),
                    item.get("n_measurements"),
                    item.get("technical_repeat_count"),
                    item.get("source_record_numbers"),
                ]
                for item in results["experimental_units"]
            ],
        )
        add_heading(doc, "Technical Source Records", level=2)
        add_table(
            doc,
            [
                "Record number",
                "Experimental unit",
                "Group",
                "Measurement",
                "Technical repeat",
                "Source location",
            ],
            [
                [
                    item.get("record_number"),
                    item.get("experimental_unit_id"),
                    item.get("group"),
                    item.get("measurement"),
                    item.get("technical_repeat_id"),
                    item.get("source_location"),
                ]
                for item in results["technical_source_records"]
            ],
        )
    else:
        params = results["curve_parameters"]
        add_table(
            doc,
            ["Parameter", "Value", "Unit or meaning"],
            [
                ["Lower asymptote L", params.get("lower_asymptote"), params.get("response_unit")],
                ["Upper asymptote U", params.get("upper_asymptote"), params.get("response_unit")],
                [
                    "Midpoint concentration C",
                    params.get("midpoint_concentration"),
                    params.get("concentration_unit"),
                ],
                ["Slope magnitude B", params.get("slope_magnitude"), "Dimensionless magnitude"],
                ["Direction", params.get("direction"), "Increasing or decreasing"],
            ],
        )
        add_heading(doc, "Curve Diagnostics", level=2)
        add_key_value_table(doc, results["diagnostics"])
        add_heading(doc, "Standard Levels", level=2)
        add_table(
            doc,
            ["Concentration", "Measurements", "Response mean", "Response SD", "Source records"],
            [
                [
                    item.get("concentration"),
                    item.get("n_measurements"),
                    item.get("response_mean"),
                    item.get("response_sd"),
                    item.get("source_record_numbers"),
                ]
                for item in results["fit_levels"]
            ],
        )
        add_heading(doc, "Residuals", level=2)
        add_table(
            doc,
            [
                "Concentration",
                "Observed response",
                "Predicted response",
                "Residual",
                "Source records",
            ],
            [
                [
                    item.get("concentration"),
                    item.get("observed"),
                    item.get("predicted"),
                    item.get("residual"),
                    item.get("source_record_numbers"),
                ]
                for item in results["residuals"]
            ],
        )
        if results.get("inverse") is not None:
            inverse = results["inverse"]
            add_heading(doc, "Research Only Sample Estimates", level=2)
            add_table(
                doc,
                [
                    "Measurement ID",
                    "Sample identifier",
                    "Response",
                    "Dilution factor",
                    "Assayed concentration",
                    "Original concentration",
                    "Status",
                    "Reason",
                ],
                [
                    [
                        item.get("measurement_id"),
                        item.get("sample_id"),
                        item.get("response"),
                        item.get("dilution_factor"),
                        item.get("concentration_in_assayed_sample"),
                        item.get("concentration_in_original_sample"),
                        item.get("status"),
                        item.get("reason"),
                    ]
                    for item in inverse["records"]
                ],
            )

    add_heading(doc, "Charts")
    for chart in data["charts"]:
        add_chart(doc, chart_dir, chart)

    add_heading(doc, "Limitations and Prohibited Interpretations")
    for item in data["limitations"]:
        doc.add_paragraph(str(item), style="List Bullet")
    if report_type == "welch_two_group":
        doc.add_paragraph(
            "The independence declaration is user supplied and not verified by software."
        )
    if report_type == "generic_grouped":
        doc.add_paragraph(
            "Measurement-row counts must not be read as independent biological sample counts."
        )

    add_heading(doc, "Traceability")
    trace = {
        "report_id": data["report_id"],
        "report_type": data["report_type"],
        "source_manifest_sha256": source.get("upstream_manifest_sha256"),
        "input_artifact_sha256": source.get("input_artifact_sha256"),
    }
    add_key_value_table(doc, trace)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


if __name__ == "__main__":
    write_report(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]))
