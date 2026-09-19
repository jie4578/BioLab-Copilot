"""Portable DOCX report rendering with python-docx."""

from __future__ import annotations

import json
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = "1F4E78"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)

_LABELS = {
    "analysis_level": "Analysis level",
    "direction": "Curve direction",
    "concentration_unit": "Concentration unit",
    "response_unit": "Response unit",
    "independent_biological_n": "Independent biological n",
    "n_measurements": "Measurement rows",
    "n_experimental_units": "Experimental units",
    "sample_sd": "Sample SD",
    "mean_difference": "Mean difference",
    "mean_difference_ci_lower": "95% CI lower",
    "mean_difference_ci_upper": "95% CI upper",
    "t_statistic": "t statistic",
    "degrees_of_freedom": "Degrees of freedom",
    "p_value": "Two-sided p value",
    "lower_asymptote": "Lower asymptote L",
    "upper_asymptote": "Upper asymptote U",
    "midpoint_concentration": "Midpoint concentration C",
    "slope_magnitude": "Slope magnitude B",
}


def _label(key: str) -> str:
    return _LABELS.get(key, key.replace("_", " ").capitalize())


def _display(value: Any) -> str:
    if value is None:
        return "Not calculated"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return format(value, ".10g")
    if isinstance(value, dict | list | tuple):
        return json.dumps(value, ensure_ascii=False, allow_nan=False, sort_keys=True)
    return str(value)


def _set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_borders(cell: Any) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), BORDER)


def _add_table(document: Any, headers: list[str], rows: list[list[Any]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        _set_cell_shading(cell, NAVY)
        _set_cell_borders(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = _display(value)
            _set_cell_borders(cell)
            if row_index % 2 == 1:
                _set_cell_shading(cell, LIGHT_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_heading(document: Any, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = BLACK


def _add_key_value_table(document: Any, values: dict[str, Any]) -> None:
    _add_table(
        document, ["Field", "Value"], [[_label(key), value] for key, value in values.items()]
    )


def _add_issues(document: Any, issues: list[dict[str, Any]]) -> None:
    if not issues:
        document.add_paragraph("No warnings or issues were recorded in the upstream artifacts.")
        return
    _add_table(
        document,
        ["Code", "Severity", "Message", "Location", "Recommended action"],
        [
            [
                item.get("code"),
                item.get("severity"),
                item.get("message"),
                item.get("location"),
                item.get("suggested_action"),
            ]
            for item in issues
        ],
    )


def _add_chart(document: Any, chart_dir: Path, relative: str) -> None:
    path = chart_dir.parent / relative
    if not path.is_file():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.2))
    caption = document.add_paragraph(relative.replace("_", " ").replace("/", " "))
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8)


def _style_document(document: Any) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = document.styles
    for style_name, size in (
        ("Normal", 10.5),
        ("Title", 20),
        ("Heading 1", 14),
        ("Heading 2", 11.5),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
    styles["Title"].font.bold = True
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.bold = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("BioLab Copilot | Deterministic scientific report")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(89, 89, 89)


def _add_results(document: Any, data: dict[str, Any]) -> None:
    results = data["results"]
    report_type = data["report_type"]
    if report_type == "generic_grouped":
        _add_table(
            document,
            ["Group", "Measurement rows", "Mean", "Median", "Minimum", "Maximum", "Sample SD"],
            [
                [
                    item.get(key)
                    for key in (
                        "group",
                        "n_measurements",
                        "mean",
                        "median",
                        "min",
                        "max",
                        "sample_sd",
                    )
                ]
                for item in results["group_statistics"]
            ],
        )
        document.add_paragraph("Independent biological n: Not calculated.")
        _add_heading(document, "Measurement Rows", level=2)
        _add_table(
            document,
            [
                "Record number",
                "Sample identifier",
                "Group",
                "Measurement",
                "Repeat type",
                "Source location",
            ],
            [
                [
                    item.get(key)
                    for key in (
                        "record_number",
                        "sample_id",
                        "group",
                        "measurement",
                        "replicate_type",
                        "source_location",
                    )
                ]
                for item in results["measurements"]
            ],
        )
    elif report_type == "welch_two_group":
        _add_table(
            document,
            ["Group", "Measurement rows", "Experimental units", "Mean", "Sample SD"],
            [
                [
                    item.get(key)
                    for key in (
                        "group",
                        "n_measurements",
                        "n_experimental_units",
                        "mean",
                        "sample_sd",
                    )
                ]
                for item in results["group_statistics"]
            ],
        )
        _add_heading(document, "Confirmed Comparison", level=2)
        _add_key_value_table(document, results["comparison"])
        _add_heading(document, "Experimental Units", level=2)
        _add_table(
            document,
            [
                "Experimental unit",
                "Group",
                "Unit value",
                "Measurement rows",
                "Technical repeat count",
                "Source records",
            ],
            [
                [
                    item.get(key)
                    for key in (
                        "experimental_unit_id",
                        "group",
                        "aggregated_value",
                        "n_measurements",
                        "technical_repeat_count",
                        "source_record_numbers",
                    )
                ]
                for item in results["experimental_units"]
            ],
        )
    else:
        parameters = results["curve_parameters"]
        _add_table(
            document,
            ["Parameter", "Value", "Unit or meaning"],
            [
                [
                    "Lower asymptote L",
                    parameters.get("lower_asymptote"),
                    parameters.get("response_unit"),
                ],
                [
                    "Upper asymptote U",
                    parameters.get("upper_asymptote"),
                    parameters.get("response_unit"),
                ],
                [
                    "Midpoint concentration C",
                    parameters.get("midpoint_concentration"),
                    parameters.get("concentration_unit"),
                ],
                ["Slope magnitude B", parameters.get("slope_magnitude"), "Dimensionless magnitude"],
                ["Direction", parameters.get("direction"), "Increasing or decreasing"],
            ],
        )
        _add_heading(document, "Curve Diagnostics", level=2)
        _add_key_value_table(document, results["diagnostics"])
        _add_heading(document, "Standard Levels", level=2)
        _add_table(
            document,
            ["Concentration", "Measurements", "Response mean", "Response SD", "Source records"],
            [
                [
                    item.get(key)
                    for key in (
                        "concentration",
                        "n_measurements",
                        "response_mean",
                        "response_sd",
                        "source_record_numbers",
                    )
                ]
                for item in results["fit_levels"]
            ],
        )
        _add_heading(document, "Residuals", level=2)
        _add_table(
            document,
            [
                "Concentration",
                "Observed response",
                "Predicted response",
                "Residual",
                "Source records",
            ],
            [
                [
                    item.get(key)
                    for key in (
                        "concentration",
                        "observed",
                        "predicted",
                        "residual",
                        "source_record_numbers",
                    )
                ]
                for item in results["residuals"]
            ],
        )
        if results.get("inverse") is not None:
            _add_heading(document, "Research-Only Sample Estimates", level=2)
            _add_table(
                document,
                [
                    "Measurement ID",
                    "Sample identifier",
                    "Response",
                    "Dilution factor",
                    "Assayed concentration",
                    "Original concentration",
                    "Status",
                    "Reason",
                ],
                [
                    [
                        item.get(key)
                        for key in (
                            "measurement_id",
                            "sample_id",
                            "response",
                            "dilution_factor",
                            "concentration_in_assayed_sample",
                            "concentration_in_original_sample",
                            "status",
                            "reason",
                        )
                    ]
                    for item in results["inverse"]["records"]
                ],
            )


def _assert_safe_docx(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "word/vbaProject.bin" in names:
            raise ValueError("Generated DOCX contains a macro project.")
        relationship_files = [name for name in names if name.endswith(".rels")]
        for relationship_file in relationship_files:
            root = ElementTree.fromstring(archive.read(relationship_file))
            for relationship in root:
                if relationship.attrib.get("TargetMode") == "External":
                    raise ValueError("Generated DOCX contains an external relationship.")


def write_docx(data: dict[str, Any], output_path: Path, chart_dir: Path) -> None:
    document = Document()
    _style_document(document)
    properties = document.core_properties
    properties.author = "BioLab Copilot"
    properties.last_modified_by = "BioLab Copilot"
    properties.title = str(data["title"])
    properties.subject = str(data["report_type"])
    fixed_timestamp = datetime(2000, 1, 1, tzinfo=UTC)
    properties.created = fixed_timestamp
    properties.modified = fixed_timestamp

    title = document.add_paragraph(style="Title")
    title.add_run(str(data["title"]))
    subtitle = document.add_paragraph()
    subtitle.add_run(f"Report type: {data['report_type']} | Report ID: {data['report_id']}")
    _add_heading(document, "Purpose and Scope")
    document.add_paragraph(str(data["scope"]))
    if data["report_type"] == "elisa_4pl":
        document.add_paragraph("Research-only; curve and quantitative range are not validated.")

    _add_heading(document, "Data and Analysis Sources")
    source = data["source"]
    _add_key_value_table(
        document,
        {
            "source_filename": source.get("source_filename"),
            "source_sha256": source.get("source_sha256"),
            "input_artifact_sha256": source.get("input_artifact_sha256"),
            "upstream_run_id": source.get("upstream_run_id"),
            "upstream_manifest_sha256": source.get("upstream_manifest_sha256"),
        },
    )
    _add_heading(document, "QC and Warning Summary")
    _add_issues(document, list(data["qc"].get("issues", [])))
    _add_heading(document, "Confirmed Methods")
    _add_key_value_table(document, dict(data["methods"]))
    _add_heading(document, "Results")
    _add_results(document, data)
    _add_heading(document, "Charts")
    for chart in data["charts"]:
        _add_chart(document, chart_dir, chart)
    _add_heading(document, "Limitations and Prohibited Interpretations")
    for item in data["limitations"]:
        document.add_paragraph(str(item), style="List Bullet")
    if data["report_type"] == "welch_two_group":
        document.add_paragraph(
            "The independence declaration is user supplied and not verified by software."
        )
    if data["report_type"] == "generic_grouped":
        document.add_paragraph(
            "Measurement-row counts must not be read as independent biological sample counts."
        )
    _add_heading(document, "Traceability")
    _add_key_value_table(
        document,
        {
            "report_id": data["report_id"],
            "report_type": data["report_type"],
            "source_manifest_sha256": source.get("upstream_manifest_sha256"),
            "input_artifact_sha256": source.get("input_artifact_sha256"),
        },
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    _assert_safe_docx(output_path)
