"""Portable, value-only XLSX report rendering with openpyxl."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "1F4E78"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"
FONT = "Arial"

_LABELS = {
    "analysis_level": "Analysis level",
    "source_filename": "Source filename",
    "source_sha256": "Source SHA-256",
    "input_artifact_sha256": "Input artifact SHA-256",
    "upstream_run_id": "Upstream run ID",
    "upstream_manifest_sha256": "Upstream manifest SHA-256",
    "independent_biological_n": "Independent biological n",
    "n_measurements": "Measurement rows",
    "n_experimental_units": "Experimental units",
    "sample_sd": "Sample SD",
    "sample_sd_reason": "Sample SD reason",
    "status": "Status",
    "reason": "Reason",
}


def _label(key: str) -> str:
    return _LABELS.get(key, key.replace("_", " "))


def _safe(value: Any) -> Any:
    """Keep report cells literal and prevent spreadsheet formula injection."""

    if isinstance(value, dict | list | tuple):
        return json.dumps(value, ensure_ascii=False, allow_nan=False, sort_keys=True)
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _style_cell(cell: Any, *, header: bool = False, alternate: bool = False) -> None:
    cell.font = Font(name=FONT, size=10, bold=header, color="FFFFFF" if header else "000000")
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin", color=BORDER),
        right=Side(style="thin", color=BORDER),
        top=Side(style="thin", color=BORDER),
        bottom=Side(style="thin", color=BORDER),
    )
    if header:
        cell.fill = PatternFill(fill_type="solid", fgColor=NAVY)
    elif alternate:
        cell.fill = PatternFill(fill_type="solid", fgColor=LIGHT_GRAY)


def _set_widths(sheet: Any, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _add_title(sheet: Any, title: str, subtitle: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = _safe(title)
    sheet["A1"].font = Font(name=FONT, size=15, bold=True, color="000000")
    sheet["A2"] = _safe(subtitle)
    sheet["A2"].font = Font(name=FONT, size=10, italic=True, color="595959")


def _add_table(
    sheet: Any,
    name: str,
    start_row: int,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[int],
) -> None:
    header_row = start_row + 1
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=_safe(header))
        _style_cell(cell, header=True)
    for row_offset, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=header_row + row_offset, column=column, value=_safe(value))
            _style_cell(cell, alternate=row_offset % 2 == 0)
    if rows:
        last_row = header_row + len(rows)
        last_column = get_column_letter(len(headers))
        table = Table(displayName=name, ref=f"A{header_row}:{last_column}{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    sheet.freeze_panes = f"A{header_row + 1}"
    _set_widths(sheet, widths)


def _add_chart(sheet: Any, chart_dir: Path, relative: str) -> None:
    image_path = chart_dir / Path(relative).name
    if not image_path.is_file():
        return
    image = XlsxImage(str(image_path))
    image.width = 600
    image.height = 330
    sheet.add_image(image, "E4")


def _add_summary(workbook: Any, data: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Report Summary")
    _add_title(sheet, str(data["title"]), f"{data['report_type']} | {data['report_id']}")
    source = data["source"]
    _add_table(
        sheet,
        "SummarySource",
        3,
        ["Field", "Value"],
        [
            ["Report type", data["report_type"]],
            ["Scope", data["scope"]],
            ["Source filename", source.get("source_filename")],
            ["Source SHA-256", source.get("source_sha256")],
            ["Upstream run ID", source.get("upstream_run_id")],
            ["Independent biological n", data["results"].get("independent_biological_n")],
        ],
        [28, 86],
    )


def _add_provenance(workbook: Any, data: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Provenance")
    _add_title(sheet, "Provenance", "Logical names and hashes only; no absolute paths.")
    rows = [[_label(key), value] for key, value in data["source"].items()]
    _add_table(sheet, "ProvenanceTable", 3, ["Field", "Value"], rows, [34, 100])


def _add_issues(workbook: Any, data: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Issues")
    _add_title(sheet, "Issues", "Warnings and limitations retained from upstream artifacts.")
    issues = data["qc"].get("issues", [])
    rows = (
        [
            [
                item.get("code"),
                item.get("severity"),
                item.get("message"),
                item.get("location"),
                item.get("suggested_action"),
            ]
            for item in issues
        ]
        if issues
        else [["", "", "No warnings or issues were recorded.", "", ""]]
    )
    _add_table(
        sheet,
        "IssuesTable",
        3,
        ["Code", "Severity", "Message", "Location", "Recommended action"],
        rows,
        [28, 14, 62, 40, 55],
    )


def _add_methods(workbook: Any, data: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Methods")
    _add_title(sheet, "Methods", "Confirmed upstream method declarations.")
    rows = [[_label(key), value] for key, value in data["methods"].items()]
    _add_table(sheet, "MethodsTable", 3, ["Method field", "Value"], rows, [40, 100])


def _add_generic(workbook: Any, data: dict[str, Any]) -> None:
    group = workbook.create_sheet("Group Statistics")
    _add_title(
        group, "Group Statistics", "Measurement rows; independent biological n is not calculated."
    )
    _add_table(
        group,
        "GroupStatisticsTable",
        3,
        ["Group", "Measurement rows", "Mean", "Median", "Minimum", "Maximum", "Sample SD"],
        [
            [
                item.get(key)
                for key in ("group", "n_measurements", "mean", "median", "min", "max", "sample_sd")
            ]
            for item in data["results"]["group_statistics"]
        ],
        [24, 18, 16, 16, 16, 16, 16],
    )
    measurements = workbook.create_sheet("Measurements")
    _add_title(
        measurements,
        "Measurements",
        "All parsed measurement rows used by the upstream descriptive result.",
    )
    _add_table(
        measurements,
        "MeasurementsTable",
        3,
        [
            "Record number",
            "Sample identifier",
            "Group",
            "Measurement",
            "Repeat type",
            "Source location",
        ],
        [
            [
                item.get(key)
                for key in (
                    "record_number",
                    "sample_id",
                    "group",
                    "measurement",
                    "replicate_type",
                    "source_location",
                )
            ]
            for item in data["results"]["measurements"]
        ],
        [16, 22, 18, 16, 18, 60],
    )


def _add_welch(workbook: Any, data: dict[str, Any]) -> None:
    group = workbook.create_sheet("Group Statistics")
    _add_title(
        group, "Group Statistics", "Experimental-unit level summaries used by Welch's method."
    )
    _add_table(
        group,
        "WelchGroupStatisticsTable",
        3,
        ["Group", "Measurement rows", "Experimental units", "Mean", "Sample SD"],
        [
            [
                item.get(key)
                for key in ("group", "n_measurements", "n_experimental_units", "mean", "sample_sd")
            ]
            for item in data["results"]["group_statistics"]
        ],
        [24, 18, 20, 18, 18],
    )
    units = workbook.create_sheet("Experimental Units")
    _add_title(
        units,
        "Experimental Units",
        "Only unit-level values are used for the inferential comparison.",
    )
    _add_table(
        units,
        "ExperimentalUnitsTable",
        3,
        [
            "Experimental unit",
            "Group",
            "Unit value",
            "Measurement rows",
            "Technical repeat count",
            "Source records",
        ],
        [
            [
                item.get("experimental_unit_id"),
                item.get("group"),
                item.get("aggregated_value"),
                item.get("n_measurements"),
                item.get("technical_repeat_count"),
                item.get("source_record_numbers"),
            ]
            for item in data["results"]["experimental_units"]
        ],
        [24, 18, 18, 18, 22, 28],
    )
    comparison = workbook.create_sheet("Comparison")
    _add_title(comparison, "Comparison", "One confirmed two-sided Welch comparison.")
    _add_table(
        comparison,
        "ComparisonTable",
        3,
        ["Field", "Value"],
        [[_label(key), value] for key, value in data["results"]["comparison"].items()],
        [42, 30],
    )
    source = workbook.create_sheet("Technical Source Records")
    _add_title(
        source,
        "Technical Source Records",
        "Source rows retained separately from experimental-unit results.",
    )
    _add_table(
        source,
        "TechnicalSourceTable",
        3,
        [
            "Record number",
            "Experimental unit",
            "Group",
            "Measurement",
            "Technical repeat",
            "Source location",
        ],
        [
            [
                item.get(key)
                for key in (
                    "record_number",
                    "experimental_unit_id",
                    "group",
                    "measurement",
                    "technical_repeat_id",
                    "source_location",
                )
            ]
            for item in data["results"]["technical_source_records"]
        ],
        [16, 24, 18, 16, 20, 60],
    )


def _add_elisa(workbook: Any, data: dict[str, Any]) -> None:
    results = data["results"]
    params = workbook.create_sheet("Curve Parameters")
    _add_title(params, "Curve Parameters", "Standard-only 4PL parameters and diagnostics.")
    _add_table(
        params,
        "CurveParametersTable",
        3,
        ["Parameter", "Value"],
        [[_label(key), value] for key, value in results["curve_parameters"].items()],
        [42, 36],
    )
    _add_table(
        params,
        "CurveDiagnosticsTable",
        13,
        ["Diagnostic", "Value"],
        [[_label(key), value] for key, value in results["diagnostics"].items()],
        [42, 60],
    )
    levels = workbook.create_sheet("Standard Levels")
    _add_title(levels, "Standard Levels", "Concentration levels used by the upstream fit.")
    _add_table(
        levels,
        "StandardLevelsTable",
        3,
        ["Concentration", "Measurements", "Response mean", "Response SD", "Source records"],
        [
            [
                item.get(key)
                for key in (
                    "concentration",
                    "n_measurements",
                    "response_mean",
                    "response_sd",
                    "source_record_numbers",
                )
            ]
            for item in results["fit_levels"]
        ],
        [18, 18, 20, 18, 30],
    )
    raw = workbook.create_sheet("Raw Standards")
    _add_title(raw, "Raw Standards", "Original standard records retained with inclusion decisions.")
    _add_table(
        raw,
        "RawStandardsTable",
        3,
        [
            "Record number",
            "Sample identifier",
            "Sample type",
            "Raw concentration",
            "Parsed concentration",
            "Raw response",
            "Parsed response",
            "Included",
            "Reason",
        ],
        [
            [
                item.get(key)
                for key in (
                    "record_number",
                    "sample_id",
                    "sample_type",
                    "raw_concentration",
                    "parsed_concentration",
                    "raw_response",
                    "parsed_response",
                    "included",
                    "exclusion_reason",
                )
            ]
            for item in results["raw_standard_records"]
        ],
        [16, 22, 18, 20, 20, 20, 20, 14, 45],
    )
    residuals = workbook.create_sheet("Residuals")
    _add_title(
        residuals,
        "Residuals",
        "Observed, predicted, and residual values from the selected fit levels.",
    )
    _add_table(
        residuals,
        "ResidualsTable",
        3,
        ["Concentration", "Observed response", "Predicted response", "Residual", "Source records"],
        [
            [
                item.get(key)
                for key in (
                    "concentration",
                    "observed",
                    "predicted",
                    "residual",
                    "source_record_numbers",
                )
            ]
            for item in results["residuals"]
        ],
        [20, 22, 22, 18, 30],
    )
    samples = workbook.create_sheet("Sample Estimates")
    _add_title(
        samples,
        "Sample Estimates",
        "Research-only per-measurement estimates; blank cells mean no estimate was calculated.",
    )
    inverse = results.get("inverse")
    sample_rows = (
        []
        if inverse is None
        else [
            [
                item.get(key)
                for key in (
                    "measurement_id",
                    "sample_id",
                    "response",
                    "dilution_factor",
                    "concentration_in_assayed_sample",
                    "concentration_in_original_sample",
                    "status",
                    "reason",
                )
            ]
            for item in inverse["records"]
        ]
    )
    _add_table(
        samples,
        "SampleEstimatesTable",
        3,
        [
            "Measurement ID",
            "Sample identifier",
            "Response",
            "Dilution factor",
            "Assayed concentration",
            "Original concentration",
            "Status",
            "Reason",
        ],
        sample_rows,
        [22, 22, 18, 18, 24, 24, 34, 60],
    )
    statuses = workbook.create_sheet("Sample Status Summary")
    _add_title(
        statuses, "Sample Status Summary", "Counts of per-measurement research-only statuses."
    )
    order = [
        "estimated_within_standard_span",
        "below_standard_span",
        "above_standard_span",
        "outside_model_domain",
        "near_asymptote_unstable",
        "numerical_failure",
    ]
    status_rows = (
        [
            [status, sum(item.get("status") == status for item in inverse["records"])]
            for status in order
        ]
        if inverse is not None
        else [["No inverse package supplied", None]]
    )
    _add_table(
        statuses, "SampleStatusTable", 3, ["Status", "Measurement rows"], status_rows, [42, 22]
    )


def _assert_safe_xlsx(path: Path) -> None:
    workbook = load_workbook(path, data_only=False, keep_links=False, read_only=False)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        raise ValueError("Generated XLSX contains an uncontrolled formula.")
        if getattr(workbook, "_external_links", []):
            raise ValueError("Generated XLSX contains external links.")
    finally:
        workbook.close()
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if any(name.startswith("xl/externalLinks/") for name in names):
            raise ValueError("Generated XLSX contains external link parts.")
        if "xl/vbaProject.bin" in names:
            raise ValueError("Generated XLSX contains a macro project.")


def write_xlsx(data: dict[str, Any], output_path: Path, chart_dir: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "BioLab Copilot"
    workbook.properties.title = str(data["title"])
    workbook.properties.subject = str(data["report_type"])
    workbook.calculation.fullCalcOnLoad = False
    _add_summary(workbook, data)
    if data["charts"]:
        _add_chart(workbook["Report Summary"], chart_dir, data["charts"][0])
    _add_provenance(workbook, data)
    _add_issues(workbook, data)
    _add_methods(workbook, data)
    if data["report_type"] == "generic_grouped":
        _add_generic(workbook, data)
    elif data["report_type"] == "welch_two_group":
        _add_welch(workbook, data)
    elif data["report_type"] == "elisa_4pl":
        _add_elisa(workbook, data)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    _assert_safe_xlsx(output_path)
