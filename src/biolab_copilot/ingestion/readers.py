"""Read-only CSV and XLSX source readers for Phase 1."""

from __future__ import annotations

import csv
import hashlib
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from biolab_copilot.contracts import InputFile, IssueSeverity, ValidationIssue


@dataclass(frozen=True)
class ReaderLimits:
    """Configurable resource limits for source files."""

    max_file_bytes: int = 50_000_000
    max_rows: int = 100_000
    max_uncompressed_bytes: int = 500_000_000
    max_zip_members: int = 10_000


@dataclass(frozen=True)
class RawRecord:
    """One logical source record before canonical parsing."""

    record_number: int
    source_file: str
    sheet_name: str | None
    source_row_number: int | None
    raw_values: dict[str, Any]
    source_locations: dict[str, str]
    formula_columns: tuple[str, ...] = ()


@dataclass(frozen=True)
class RawTable:
    """A source table and reader-level diagnostics."""

    input_file: InputFile
    headers: tuple[str, ...]
    records: tuple[RawRecord, ...]
    issues: tuple[ValidationIssue, ...]
    format: str
    sheet_name: str | None


class InputReadError(Exception):
    """A source could not be read into a table."""

    def __init__(self, input_file: InputFile | None, issues: list[ValidationIssue]) -> None:
        self.input_file = input_file
        self.issues = issues
        super().__init__("; ".join(issue.message for issue in issues))


def sha256_file(path: Path) -> str:
    """Return the SHA-256 digest of the exact source bytes."""

    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _safe_raw_value(value: Any, limit: int = 200) -> str | None:
    if value is None:
        return None
    text = str(value)
    return text if len(text) <= limit else f"{text[: limit - 3]}..."


def _issue(
    *,
    code: str,
    severity: IssueSeverity,
    message: str,
    location: str,
    suggested_action: str,
    source_file: str | None = None,
    sheet_name: str | None = None,
    record_number: int | None = None,
    source_row_number: int | None = None,
    field: str | None = None,
    raw_value: Any = None,
) -> ValidationIssue:
    return ValidationIssue(
        code=code,
        severity=severity,
        message=message,
        location=location,
        suggested_action=suggested_action,
        source_file=source_file,
        sheet_name=sheet_name,
        record_number=record_number,
        source_row_number=source_row_number,
        field=field,
        raw_value=_safe_raw_value(raw_value),
    )


def _input_file(path: Path, digest: str) -> InputFile:
    return InputFile(path=str(path.resolve()), sha256=digest, size_bytes=path.stat().st_size)


def _header_issue_list(
    headers: list[str], source_file: str, sheet_name: str | None
) -> list[ValidationIssue]:
    issues: list[ValidationIssue] = []
    if not headers:
        issues.append(
            _issue(
                code="MISSING_HEADER_ROW",
                severity=IssueSeverity.BLOCKING,
                message="The source does not contain a header row.",
                location=f"file={source_file}; sheet={sheet_name or '-'}; header",
                suggested_action=(
                    "Provide a non-empty header row and submit an explicit column mapping."
                ),
                source_file=source_file,
                sheet_name=sheet_name,
            )
        )
        return issues
    seen: set[str] = set()
    for index, header in enumerate(headers, start=1):
        if header == "":
            issues.append(
                _issue(
                    code="MISSING_SOURCE_COLUMN_NAME",
                    severity=IssueSeverity.BLOCKING,
                    message=f"Source column {index} has an empty header.",
                    location=(
                        f"file={source_file}; sheet={sheet_name or '-'}; header_column={index}"
                    ),
                    suggested_action="Give every source column a unique non-empty header.",
                    source_file=source_file,
                    sheet_name=sheet_name,
                    raw_value=header,
                )
            )
        if header in seen:
            issues.append(
                _issue(
                    code="DUPLICATE_SOURCE_COLUMN",
                    severity=IssueSeverity.BLOCKING,
                    message=f"Source column header {header!r} appears more than once.",
                    location=f"file={source_file}; sheet={sheet_name or '-'}; header={header!r}",
                    suggested_action="Rename duplicate source headers before importing.",
                    source_file=source_file,
                    sheet_name=sheet_name,
                    raw_value=header,
                )
            )
        seen.add(header)
    return issues


def _check_limits(path: Path, limits: ReaderLimits, digest: str) -> InputFile:
    size = path.stat().st_size
    input_file = _input_file(path, digest)
    if size > limits.max_file_bytes:
        raise InputReadError(
            input_file,
            [
                _issue(
                    code="FILE_SIZE_LIMIT_EXCEEDED",
                    severity=IssueSeverity.BLOCKING,
                    message=f"The file is {size} bytes, above the configured limit.",
                    location=f"file={path.name}",
                    suggested_action=(
                        "Use a smaller file or explicitly increase the file-size limit."
                    ),
                    source_file=path.name,
                )
            ],
        )
    return input_file


def _record_source_location(source_file: str, record_number: int, column: str) -> str:
    return f"file={source_file}; record={record_number}; column={column!r}"


def _row_to_record(
    *,
    row: list[Any],
    headers: list[str],
    record_number: int,
    source_file: str,
    sheet_name: str | None,
    source_row_number: int | None,
    formula_columns: tuple[str, ...] = (),
) -> RawRecord:
    values: dict[str, Any] = {}
    locations: dict[str, str] = {}
    for index, header in enumerate(headers):
        value = row[index] if index < len(row) else None
        values[header] = value
        if sheet_name is None:
            locations[header] = _record_source_location(source_file, record_number, header)
        else:
            cell = f"{get_column_letter(index + 1)}{source_row_number}"
            locations[header] = f"file={source_file}; sheet={sheet_name!r}; cell={cell}"
    for extra_index, value in enumerate(row[len(headers) :], start=1):
        extra_name = f"__extra_column_{extra_index}"
        values[extra_name] = value
        locations[extra_name] = _record_source_location(source_file, record_number, extra_name)
    return RawRecord(
        record_number=record_number,
        source_file=source_file,
        sheet_name=sheet_name,
        source_row_number=source_row_number,
        raw_values=values,
        source_locations=locations,
        formula_columns=formula_columns,
    )


def _read_csv(
    path: Path,
    input_file: InputFile,
    *,
    encoding: str,
    delimiter: str,
    limits: ReaderLimits,
) -> RawTable:
    source_file = path.name
    issues: list[ValidationIssue] = []
    records: list[RawRecord] = []
    if path.stat().st_size == 0:
        raise InputReadError(
            input_file,
            [
                _issue(
                    code="EMPTY_FILE",
                    severity=IssueSeverity.BLOCKING,
                    message="The CSV file is empty.",
                    location=f"file={source_file}",
                    suggested_action="Provide a CSV file with a header row and data records.",
                    source_file=source_file,
                )
            ],
        )
    try:
        with path.open("r", encoding=encoding, newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter, strict=True)
            try:
                headers = next(reader)
            except StopIteration:
                headers = []
            issues.extend(_header_issue_list(headers, source_file, None))
            for record_number, row in enumerate(reader, start=1):
                if record_number > limits.max_rows:
                    issues.append(
                        _issue(
                            code="ROW_LIMIT_EXCEEDED",
                            severity=IssueSeverity.BLOCKING,
                            message=f"The CSV contains more than {limits.max_rows} data records.",
                            location=f"file={source_file}; record={record_number}",
                            suggested_action=(
                                "Use a smaller file or explicitly increase the row limit."
                            ),
                            source_file=source_file,
                            record_number=record_number,
                        )
                    )
                    break
                if not row:
                    continue
                if len(row) != len(headers):
                    issues.append(
                        _issue(
                            code="ROW_WIDTH_MISMATCH",
                            severity=IssueSeverity.ERROR,
                            message=(
                                f"Record {record_number} has {len(row)} fields; "
                                f"expected {len(headers)}."
                            ),
                            location=f"file={source_file}; record={record_number}",
                            suggested_action=(
                                "Repair the CSV row while preserving the original source file."
                            ),
                            source_file=source_file,
                            record_number=record_number,
                        )
                    )
                if (
                    headers
                    and len(row) == len(headers)
                    and all(
                        str(value).strip().lower() == str(header).strip().lower()
                        for value, header in zip(row, headers, strict=True)
                    )
                ):
                    issues.append(
                        _issue(
                            code="DUPLICATE_HEADER_ROW",
                            severity=IssueSeverity.ERROR,
                            message=f"Record {record_number} repeats the header row.",
                            location=f"file={source_file}; record={record_number}",
                            suggested_action=(
                                "Remove the repeated header from a derived copy, then re-import "
                                "the original source after review."
                            ),
                            source_file=source_file,
                            record_number=record_number,
                        )
                    )
                records.append(
                    _row_to_record(
                        row=row,
                        headers=headers,
                        record_number=record_number,
                        source_file=source_file,
                        sheet_name=None,
                        source_row_number=None,
                    )
                )
    except UnicodeDecodeError as exc:
        raise InputReadError(
            input_file,
            [
                _issue(
                    code="CSV_ENCODING_ERROR",
                    severity=IssueSeverity.BLOCKING,
                    message=(
                        "The CSV could not be decoded with the explicitly selected encoding: "
                        f"{exc.encoding}."
                    ),
                    location=f"file={source_file}",
                    suggested_action=(
                        "Retry with an explicit correct encoding; no encoding is guessed "
                        "automatically."
                    ),
                    source_file=source_file,
                )
            ],
        ) from exc
    except LookupError as exc:
        raise InputReadError(
            input_file,
            [
                _issue(
                    code="UNSUPPORTED_ENCODING",
                    severity=IssueSeverity.BLOCKING,
                    message=f"The requested text encoding is unavailable: {encoding}.",
                    location=f"file={source_file}",
                    suggested_action="Use a Python-supported encoding explicitly.",
                    source_file=source_file,
                )
            ],
        ) from exc
    except csv.Error as exc:
        raise InputReadError(
            input_file,
            [
                _issue(
                    code="CSV_MALFORMED_QUOTING",
                    severity=IssueSeverity.BLOCKING,
                    message=f"The CSV contains malformed quoting: {exc}.",
                    location=f"file={source_file}",
                    suggested_action=(
                        "Repair quoting in a reviewed derived copy; preserve the original "
                        "source file."
                    ),
                    source_file=source_file,
                )
            ],
        ) from exc
    return RawTable(input_file, tuple(headers), tuple(records), tuple(issues), "csv", None)


def _check_xlsx_zip_limits(path: Path, input_file: InputFile, limits: ReaderLimits) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            uncompressed_size = sum(info.file_size for info in members)
    except zipfile.BadZipFile as exc:
        raise InputReadError(
            input_file,
            [
                _issue(
                    code="CORRUPT_XLSX",
                    severity=IssueSeverity.BLOCKING,
                    message="The .xlsx file is not a valid ZIP-based workbook.",
                    location=f"file={path.name}",
                    suggested_action="Provide an uncorrupted .xlsx file.",
                    source_file=path.name,
                )
            ],
        ) from exc
    if len(members) > limits.max_zip_members or uncompressed_size > limits.max_uncompressed_bytes:
        raise InputReadError(
            input_file,
            [
                _issue(
                    code="XLSX_UNCOMPRESSED_LIMIT_EXCEEDED",
                    severity=IssueSeverity.BLOCKING,
                    message=(
                        "The workbook exceeds the configured ZIP member or uncompressed-size limit."
                    ),
                    location=f"file={path.name}",
                    suggested_action=(
                        "Use a smaller workbook or explicitly increase the safe resource limits."
                    ),
                    source_file=path.name,
                )
            ],
        )


def _xlsx_headers(cells: tuple[Any, ...]) -> list[str]:
    headers: list[str] = []
    for cell in cells:
        value = cell.value
        headers.append("" if value is None else str(value))
    return headers


def _read_xlsx(
    path: Path, input_file: InputFile, *, sheet_name: str | None, limits: ReaderLimits
) -> RawTable:
    source_file = path.name
    _check_xlsx_zip_limits(path, input_file, limits)
    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=False,
            keep_links=False,
            keep_vba=False,
        )
    except (InvalidFileException, ValueError, OSError, zipfile.BadZipFile) as exc:
        raise InputReadError(
            input_file,
            [
                _issue(
                    code="CORRUPT_XLSX",
                    severity=IssueSeverity.BLOCKING,
                    message=f"The workbook could not be opened safely: {type(exc).__name__}.",
                    location=f"file={source_file}",
                    suggested_action="Provide an uncorrupted, unencrypted .xlsx file.",
                    source_file=source_file,
                )
            ],
        ) from exc
    try:
        sheet_names = list(workbook.sheetnames)
        if sheet_name is None and len(sheet_names) > 1:
            raise InputReadError(
                input_file,
                [
                    _issue(
                        code="SHEET_SELECTION_REQUIRED",
                        severity=IssueSeverity.BLOCKING,
                        message=f"The workbook contains multiple sheets: {', '.join(sheet_names)}.",
                        location=f"file={source_file}; workbook",
                        suggested_action=(
                            "List the sheets and select one explicitly; sheets are never "
                            "auto-concatenated."
                        ),
                        source_file=source_file,
                        raw_value=", ".join(sheet_names),
                    )
                ],
            )
        selected_sheet = sheet_name or sheet_names[0]
        if selected_sheet not in sheet_names:
            raise InputReadError(
                input_file,
                [
                    _issue(
                        code="SHEET_NOT_FOUND",
                        severity=IssueSeverity.BLOCKING,
                        message=f"Worksheet {selected_sheet!r} does not exist.",
                        location=f"file={source_file}; workbook",
                        suggested_action=(
                            f"Choose one of the available worksheets: {', '.join(sheet_names)}."
                        ),
                        source_file=source_file,
                        sheet_name=selected_sheet,
                    )
                ],
            )
        worksheet = workbook[selected_sheet]
        row_iterator: Iterator[tuple[Any, ...]] = worksheet.iter_rows()
        try:
            header_cells = next(row_iterator)
        except StopIteration:
            header_cells = ()
        headers = _xlsx_headers(header_cells)
        issues = _header_issue_list(headers, source_file, selected_sheet)
        records: list[RawRecord] = []
        for row_number, cells in enumerate(row_iterator, start=2):
            values = [cell.value for cell in cells]
            if not values or all(value is None for value in values):
                continue
            record_number = len(records) + 1
            if record_number > limits.max_rows:
                issues.append(
                    _issue(
                        code="ROW_LIMIT_EXCEEDED",
                        severity=IssueSeverity.BLOCKING,
                        message=f"The worksheet contains more than {limits.max_rows} data records.",
                        location=f"file={source_file}; sheet={selected_sheet!r}; row={row_number}",
                        suggested_action=(
                            "Use a smaller workbook or explicitly increase the row limit."
                        ),
                        source_file=source_file,
                        sheet_name=selected_sheet,
                        source_row_number=row_number,
                    )
                )
                break
            if (
                headers
                and len(values) == len(headers)
                and all(
                    str(value).strip().lower() == str(header).strip().lower()
                    for value, header in zip(values, headers, strict=True)
                )
            ):
                issues.append(
                    _issue(
                        code="DUPLICATE_HEADER_ROW",
                        severity=IssueSeverity.ERROR,
                        message=f"Worksheet row {row_number} repeats the header row.",
                        location=f"file={source_file}; sheet={selected_sheet!r}; row={row_number}",
                        suggested_action=(
                            "Remove the repeated header from a reviewed derived copy, then "
                            "re-import the original source after review."
                        ),
                        source_file=source_file,
                        sheet_name=selected_sheet,
                        record_number=record_number,
                        source_row_number=row_number,
                    )
                )
            formula_columns = tuple(
                header
                for header, cell in zip(headers, cells, strict=False)
                if cell.data_type == "f"
            )
            records.append(
                _row_to_record(
                    row=values,
                    headers=headers,
                    record_number=record_number,
                    source_file=source_file,
                    sheet_name=selected_sheet,
                    source_row_number=row_number,
                    formula_columns=formula_columns,
                )
            )
        return RawTable(
            input_file, tuple(headers), tuple(records), tuple(issues), "xlsx", selected_sheet
        )
    finally:
        workbook.close()


def read_source(
    path: str | Path,
    *,
    sheet_name: str | None = None,
    encoding: str = "utf-8-sig",
    delimiter: str = ",",
    limits: ReaderLimits | None = None,
) -> RawTable:
    """Read a CSV or XLSX without changing its bytes or guessing its schema."""

    source_path = Path(path).resolve()
    if not source_path.is_file():
        raise InputReadError(
            None,
            [
                _issue(
                    code="INPUT_FILE_NOT_FOUND",
                    severity=IssueSeverity.BLOCKING,
                    message=f"Input file does not exist: {source_path.name}.",
                    location=f"file={source_path}",
                    suggested_action="Provide an existing local CSV or XLSX file.",
                    source_file=source_path.name,
                )
            ],
        )
    if len(delimiter) != 1:
        raise InputReadError(
            None,
            [
                _issue(
                    code="INVALID_DELIMITER",
                    severity=IssueSeverity.BLOCKING,
                    message="The delimiter must be exactly one character.",
                    location=f"file={source_path.name}",
                    suggested_action="Pass one explicit delimiter character.",
                    source_file=source_path.name,
                )
            ],
        )
    configured_limits = limits or ReaderLimits()
    before_hash = sha256_file(source_path)
    input_file = _check_limits(source_path, configured_limits, before_hash)
    suffix = source_path.suffix.lower()
    try:
        if suffix == ".csv":
            table = _read_csv(
                source_path,
                input_file,
                encoding=encoding,
                delimiter=delimiter,
                limits=configured_limits,
            )
        elif suffix == ".xlsx":
            table = _read_xlsx(
                source_path,
                input_file,
                sheet_name=sheet_name,
                limits=configured_limits,
            )
        else:
            raise InputReadError(
                input_file,
                [
                    _issue(
                        code="UNSUPPORTED_FILE_FORMAT",
                        severity=IssueSeverity.BLOCKING,
                        message=(
                            "Only .csv and .xlsx are supported; received "
                            f"{source_path.suffix or 'no extension'}."
                        ),
                        location=f"file={source_path.name}",
                        suggested_action=(
                            "Convert the source in a reviewed workflow or provide a supported "
                            "file type."
                        ),
                        source_file=source_path.name,
                    )
                ],
            )
    except InputReadError as exc:
        after_hash = sha256_file(source_path)
        if after_hash != before_hash:
            exc.issues.append(
                _issue(
                    code="INPUT_HASH_CHANGED_DURING_READ",
                    severity=IssueSeverity.BLOCKING,
                    message="The input bytes changed between the pre-read and post-read hash.",
                    location=f"file={source_path.name}",
                    suggested_action="Use a stable copy of the source and rerun the import.",
                    source_file=source_path.name,
                )
            )
        raise
    after_hash = sha256_file(source_path)
    if after_hash != before_hash:
        table = RawTable(
            table.input_file,
            table.headers,
            table.records,
            table.issues
            + (
                _issue(
                    code="INPUT_HASH_CHANGED_DURING_READ",
                    severity=IssueSeverity.BLOCKING,
                    message="The input bytes changed between the pre-read and post-read hash.",
                    location=f"file={source_path.name}",
                    suggested_action="Use a stable copy of the source and rerun the import.",
                    source_file=source_path.name,
                ),
            ),
            table.format,
            table.sheet_name,
        )
    return table


def list_xlsx_sheets(path: str | Path, *, limits: ReaderLimits | None = None) -> list[str]:
    """List worksheet names without selecting or concatenating a worksheet."""

    source_path = Path(path).resolve()
    if source_path.suffix.lower() != ".xlsx":
        raise InputReadError(
            None,
            [
                _issue(
                    code="UNSUPPORTED_FILE_FORMAT",
                    severity=IssueSeverity.BLOCKING,
                    message="Worksheet listing supports only .xlsx files.",
                    location=f"file={source_path.name}",
                    suggested_action="Provide an .xlsx file.",
                    source_file=source_path.name,
                )
            ],
        )
    digest = sha256_file(source_path)
    input_file = _check_limits(source_path, limits or ReaderLimits(), digest)
    _check_xlsx_zip_limits(source_path, input_file, limits or ReaderLimits())
    try:
        workbook = load_workbook(
            filename=source_path,
            read_only=True,
            data_only=False,
            keep_links=False,
            keep_vba=False,
        )
    except (InvalidFileException, ValueError, OSError, zipfile.BadZipFile) as exc:
        raise InputReadError(
            input_file,
            [
                _issue(
                    code="CORRUPT_XLSX",
                    severity=IssueSeverity.BLOCKING,
                    message="The workbook could not be opened safely.",
                    location=f"file={source_path.name}",
                    suggested_action="Provide an uncorrupted, unencrypted .xlsx file.",
                    source_file=source_path.name,
                )
            ],
        ) from exc
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


__all__ = [
    "InputReadError",
    "RawRecord",
    "RawTable",
    "ReaderLimits",
    "list_xlsx_sheets",
    "read_source",
    "sha256_file",
]
