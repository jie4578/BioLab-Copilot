"""Phase 1 ingestion and structural QC tests using synthetic data only."""

from __future__ import annotations

import hashlib
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from openpyxl import Workbook

from biolab_copilot.contracts import IssueSeverity
from biolab_copilot.ingestion import InputReadError, ReaderLimits, read_source, sha256_file
from biolab_copilot.profiling import profile_and_validate, validate_mapping

PROJECT_ROOT = Path(__file__).resolve().parents[2]


def _case_dir() -> TemporaryDirectory[str]:
    return TemporaryDirectory(prefix="phase1-", dir=PROJECT_ROOT / "tests")


def _generic_mapping() -> dict[str, str]:
    return {
        "sample_id": "sample_id",
        "group": "group",
        "measurement": "measurement",
        "replicate": "replicate_id",
        "unit": "unit",
    }


def _elisa_mapping() -> dict[str, str]:
    return {
        "sample_id": "sample_id",
        "sample_type": "sample_type",
        "standard_concentration": "standard_concentration",
        "measurement": "measurement",
        "unit": "concentration_unit",
        "replicate_type": "replicate_type",
    }


def _write(path: Path, content: bytes | str) -> str:
    if isinstance(content, str):
        path.write_text(content, encoding="utf-8")
        encoded = content.encode("utf-8")
    else:
        path.write_bytes(content)
        encoded = content
    return hashlib.sha256(encoded).hexdigest()


def test_csv_bom_unicode_quoted_newline_identifiers_and_zero_are_preserved() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "unicode.csv"
        content = (
            "样本ID,组别,测量,单位,备注\n"
            '001,control,0,µg/mL,"a,b\nline"\n'
            "NA,control,1,µg/mL,plain\n"
        ).encode("utf-8-sig")
        expected_hash = _write(path, content)
        mapping = {"样本ID": "sample_id", "组别": "group", "测量": "measurement", "单位": "unit"}

        before_hash = sha256_file(path)
        table = read_source(path)
        result = profile_and_validate(table, "generic_grouped", mapping)
        after_hash = sha256_file(path)

        assert before_hash == expected_hash == after_hash == table.input_file.sha256
        assert result.analysis_ready is True
        assert len(result.records) == 2
        assert result.records[0].parsed_values["sample_id"] == "001"
        assert result.records[0].parsed_values["measurement"] == 0.0
        assert result.records[1].parsed_values["sample_id"] == "NA"
        assert result.records[0].raw_values["备注"] == "a,b\nline"
        assert table.records[0].source_locations["备注"] == (
            "file=unicode.csv; record=1; column='备注'"
        )
        assert any(issue.code == "REPLICATE_TYPE_UNKNOWN" for issue in result.validation_issues)
        assert result.model_validate_json(result.model_dump_json()) == result
        assert "field_profiles" in result.dataset_profile.model_json_schema()["properties"]


def test_generic_invalid_nonfinite_missing_duplicate_and_unit_mix_are_visible() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "qc.csv"
        _write(
            path,
            (
                "sample_id,group,measurement,unit\n"
                "S1,control,1,ng/mL\n"
                "S1,control,1,ng/mL\n"
                "S2,,not-a-number,ng/mL\n"
                "S3,treatment,NaN,µg/mL\n"
            ),
        )

        result = profile_and_validate(
            read_source(path),
            "generic_grouped",
            {
                "sample_id": "sample_id",
                "group": "group",
                "measurement": "measurement",
                "unit": "unit",
            },
        )
        codes = {issue.code for issue in result.validation_issues}

        assert result.analysis_ready is False
        assert len(result.records) == 4
        assert result.dataset_profile.duplicate_row_count == 1
        assert {
            "DUPLICATE_COMPLETE_RECORD",
            "EMPTY_GROUP",
            "INVALID_NUMERIC",
            "NON_FINITE_NUMERIC",
            "MIXED_UNITS",
        } <= codes
        assert result.dataset_profile.error_record_count >= 2


def test_elisa_standard_concentration_is_conditional_and_unknown_concentration_stays_missing() -> (
    None
):
    with _case_dir() as temp:
        path = Path(temp) / "elisa.csv"
        _write(
            path,
            (
                "sample_id,sample_type,standard_concentration,measurement,unit,replicate_type\n"
                "STD-1,standard,0.1,0.08,ng/mL,technical\n"
                "STD-2,standard,1.0,0.55,ng/mL,technical\n"
                "SAMPLE-1,sample,,0.30,ng/mL,biological\n"
            ),
        )

        result = profile_and_validate(read_source(path), "elisa_standard_curve", _elisa_mapping())

        assert result.analysis_ready is True
        assert len(result.records) == 3
        assert result.records[2].parsed_values["standard_concentration"] is None
        assert result.records[2].parsed_values["replicate_type"] == "biological"
        assert not any(
            issue.code == "STANDARD_CONCENTRATION_REQUIRED" for issue in result.validation_issues
        )


def test_normal_elisa_xlsx_preserves_sheet_and_numeric_fields() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "elisa.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Plate 1"
        sheet.append(
            [
                "sample_id",
                "sample_type",
                "standard_concentration",
                "measurement",
                "unit",
                "replicate_type",
            ]
        )
        sheet.append(["STD-1", "standard", 0.1, 0.08, "ng/mL", "technical"])
        sheet.append(["SAMPLE-1", "sample", None, 0.30, "ng/mL", "biological"])
        workbook.save(path)

        result = profile_and_validate(
            read_source(path, sheet_name="Plate 1"),
            "elisa_standard_curve",
            _elisa_mapping(),
        )

        assert result.analysis_ready is True
        assert result.dataset_profile.source_sheet_name == "Plate 1"
        assert result.records[0].parsed_values["standard_concentration"] == 0.1
        assert result.records[1].parsed_values["standard_concentration"] is None


def test_elisa_standard_without_concentration_is_blocking_and_no_fit_is_attempted() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "elisa_missing_standard.csv"
        _write(
            path,
            (
                "sample_id,sample_type,standard_concentration,measurement\n"
                "STD-1,standard,,0.08\n"
                "SAMPLE-1,sample,,0.30\n"
            ),
        )

        result = profile_and_validate(
            read_source(path),
            "elisa_standard_curve",
            {
                "sample_id": "sample_id",
                "sample_type": "sample_type",
                "standard_concentration": "standard_concentration",
                "measurement": "measurement",
            },
        )

        assert result.analysis_ready is False
        assert any(
            issue.code == "STANDARD_CONCENTRATION_REQUIRED"
            and issue.severity == IssueSeverity.BLOCKING
            for issue in result.validation_issues
        )


def test_xlsx_requires_sheet_selection_and_rejects_mapped_formula_cells() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "multi.xlsx"
        workbook = Workbook()
        first = workbook.active
        first.title = "Data"
        first.append(["sample_id", "group", "measurement"])
        first.append(["001", "control", "=1+1"])
        workbook.create_sheet("Other").append(["unrelated"])
        workbook.save(path)
        before_hash = sha256_file(path)

        with pytest.raises(InputReadError, match="multiple sheets") as error:
            read_source(path)
        assert error.value.issues[0].code == "SHEET_SELECTION_REQUIRED"

        table = read_source(path, sheet_name="Data")
        result = profile_and_validate(
            table,
            "generic_grouped",
            {"sample_id": "sample_id", "group": "group", "measurement": "measurement"},
        )

        assert result.analysis_ready is False
        assert result.dataset_profile.source_sheet_name == "Data"
        assert result.records[0].source_row_number == 2
        assert result.records[0].raw_values["measurement"] == "=1+1"
        assert any(
            issue.code == "FORMULA_CELL_IN_MAPPED_FIELD" for issue in result.validation_issues
        )
        assert sha256_file(path) == before_hash


def test_xlsx_numeric_identifier_is_not_padded_and_warns() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "numeric_id.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Measurements"
        sheet.append(["sample_id", "group", "measurement"])
        sheet.append([1, "control", 0])
        workbook.save(path)

        result = profile_and_validate(
            read_source(path, sheet_name="Measurements"),
            "generic_grouped",
            {"sample_id": "sample_id", "group": "group", "measurement": "measurement"},
        )

        assert result.analysis_ready is True
        assert result.records[0].parsed_values["sample_id"] == "1"
        assert any(issue.code == "NUMERIC_IDENTIFIER_CELL" for issue in result.validation_issues)


def test_mapping_errors_are_explicit_and_do_not_discard_source_records() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "mapping.csv"
        _write(path, b"sample_id,group,measurement\nS1,control,1\n")
        table = read_source(path)
        issues = validate_mapping(
            table,
            "generic_grouped",
            {"missing": "sample_id", "group": "sample_id"},
        )
        codes = {issue.code for issue in issues}
        assert {
            "MAPPED_SOURCE_COLUMN_MISSING",
            "DUPLICATE_TARGET_FIELD_MAPPING",
            "MISSING_REQUIRED_FIELD_MAPPING",
        } <= codes

        result = profile_and_validate(table, "generic_grouped", {"missing": "sample_id"})
        assert len(result.records) == 1
        assert result.analysis_ready is False


def test_malformed_csv_and_limits_are_blocking() -> None:
    with _case_dir() as temp:
        malformed = Path(temp) / "malformed.csv"
        _write(malformed, b'sample_id,group,measurement\nS1,control,"unterminated\n')
        with pytest.raises(InputReadError) as error:
            read_source(malformed)
        assert error.value.issues[0].code == "CSV_MALFORMED_QUOTING"

        oversized = Path(temp) / "oversized.csv"
        _write(oversized, b"sample_id,group,measurement\nS1,control,1\n")
        with pytest.raises(InputReadError) as limit_error:
            read_source(oversized, limits=ReaderLimits(max_file_bytes=5))
        assert limit_error.value.issues[0].code == "FILE_SIZE_LIMIT_EXCEEDED"


def test_csv_row_limit_is_enforced_and_records_are_not_presented_as_complete() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "row_limit.csv"
        _write(
            path,
            "sample_id,group,measurement\nS1,control,1\nS2,control,2\n",
        )

        table = read_source(path, limits=ReaderLimits(max_rows=1))
        assert len(table.records) == 1
        assert any(issue.code == "ROW_LIMIT_EXCEEDED" for issue in table.issues)
        result = profile_and_validate(
            table,
            "generic_grouped",
            {"sample_id": "sample_id", "group": "group", "measurement": "measurement"},
        )
        assert result.analysis_ready is False
        assert any(issue.code == "ROW_LIMIT_EXCEEDED" for issue in result.validation_issues)


def test_xlsx_uncompressed_limit_is_enforced_before_workbook_processing() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "resource_limit.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(["sample_id", "group", "measurement"])
        sheet.append(["S1", "control", 1])
        workbook.save(path)
        before_hash = sha256_file(path)

        with pytest.raises(InputReadError) as error:
            read_source(
                path,
                sheet_name="Data",
                limits=ReaderLimits(max_uncompressed_bytes=1),
            )

        assert error.value.issues[0].code == "XLSX_UNCOMPRESSED_LIMIT_EXCEEDED"
        assert sha256_file(path) == before_hash


def test_empty_file_and_wrong_default_encoding_are_blocking() -> None:
    with _case_dir() as temp:
        empty = Path(temp) / "empty.csv"
        empty.write_bytes(b"")
        with pytest.raises(InputReadError) as empty_error:
            read_source(empty)
        assert empty_error.value.issues[0].code == "EMPTY_FILE"

        latin1 = Path(temp) / "latin1.csv"
        latin1.write_bytes("sample_id,group,measurement\nS1,café,1\n".encode("cp1252"))
        with pytest.raises(InputReadError) as encoding_error:
            read_source(latin1)
        assert encoding_error.value.issues[0].code == "CSV_ENCODING_ERROR"
        table = read_source(latin1, encoding="cp1252")
        assert table.records[0].raw_values["group"] == "café"


def test_corrupt_xlsx_is_blocking() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "corrupt.xlsx"
        path.write_bytes(b"not-a-zip-workbook")
        with pytest.raises(InputReadError) as error:
            read_source(path)
        assert error.value.issues[0].code == "CORRUPT_XLSX"


def test_repeated_header_and_na_are_not_silently_cleaned() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "repeat.csv"
        _write(
            path,
            "sample_id,group,measurement\nsample_id,group,measurement\nNA,control,1\n",
        )
        result = profile_and_validate(
            read_source(path),
            "generic_grouped",
            {"sample_id": "sample_id", "group": "group", "measurement": "measurement"},
        )
        assert len(result.records) == 2
        assert result.records[1].parsed_values["sample_id"] == "NA"
        assert any(issue.code == "DUPLICATE_HEADER_ROW" for issue in result.validation_issues)


def test_unsupported_file_format_is_rejected() -> None:
    with _case_dir() as temp:
        path = Path(temp) / "legacy.xls"
        _write(path, b"not an xls file")
        with pytest.raises(InputReadError) as error:
            read_source(path)
        assert error.value.issues[0].code == "UNSUPPORTED_FILE_FORMAT"
